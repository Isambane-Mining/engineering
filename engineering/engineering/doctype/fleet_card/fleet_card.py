# Copyright (c) 2026, BuFf0k and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils import add_days, date_diff, getdate, nowdate

from engineering.controllers.fleet_compliance import get_expiring_threshold_days


def _current_fleet_card(asset):
	"""The most recently issued non-cancelled Fleet Card for this Asset.
	allocation_type="Asset" is explicit (not just implied by matching
	value) - fleet_number is a Dynamic Link now, so an Asset name and an
	Employee id are different ID spaces that could theoretically collide."""
	rows = frappe.get_all(
		"Fleet Card",
		filters={"allocation_type": "Asset", "fleet_number": asset, "docstatus": ["<", 2]},
		fields=["name"],
		order_by="issue_date desc, docstatus desc",
		limit_page_length=1,
	)
	return rows[0].name if rows else None


def _expiry_status(expiry_date, threshold_days=None):
	"""Active/Expiring/Expired from Expiry Date alone — the date-driven part
	of status, shared by before_submit's initial set and the daily
	refresh_all_statuses() sweep. Never returns Superseded/Cancelled - those
	are event-driven (reissue/cancel), not date-driven, and are set directly
	at the point they happen instead."""
	if not expiry_date:
		return "Active"

	if threshold_days is None:
		threshold_days = get_expiring_threshold_days()

	today = getdate(nowdate())
	expiry_date = getdate(expiry_date)

	if expiry_date < today:
		return "Expired"
	if expiry_date <= add_days(today, threshold_days):
		return "Expiring"
	return "Active"


def refresh_all_statuses():
	"""Scheduled daily job (see hooks.py) - status is a stored field (so the
	list view can filter/report on it), but Active/Expiring/Expired are
	inherently date-driven: nothing else touches a card merely because a day
	has passed and it crossed the expiring threshold. This keeps that part
	of status from going stale by more than a day. Superseded and Cancelled
	are left alone here - those are set directly at the moment of reissue/
	cancellation (see FleetCard.on_submit/on_cancel), not on a timer."""
	threshold_days = get_expiring_threshold_days()

	cards = frappe.get_all(
		"Fleet Card",
		filters={"docstatus": 1, "status": ["in", ["Active", "Expiring", "Expired"]]},
		fields=["name", "expiry_date", "status"],
	)

	updated = 0
	for card in cards:
		new_status = _expiry_status(card.expiry_date, threshold_days)
		if new_status != card.status:
			frappe.db.set_value("Fleet Card", card.name, "status", new_status, update_modified=False)
			updated += 1

	if updated:
		frappe.db.commit()

	return updated


def _capabilities_label(doc):
	parts = []
	if doc.covers_fuel:
		parts.append("FUEL")
	if doc.covers_toll:
		parts.append("TOLL")
	if doc.covers_oil:
		parts.append("OIL")
	return "/".join(parts)


@frappe.whitelist()
def export_fleet_card_assignments_xlsx():
	"""XLSX export matching the layout of the manually-maintained "Fleet
	Cards" control sheet this replaces (REG NR / FLEET NR / Driver / Bank /
	Card Nr / Card Pin / Exp Date / Card Capabilities / Card Limit /
	Collected / Status) — the current card-to-vehicle/driver assignment
	only, not that sheet's repeated reissue-history column blocks further
	to the right (Fleet Card's own reissue history lives as separate
	documents now, linked via Previous Fleet Card, not extra columns).

	Includes every submitted card regardless of Status - Superseded and
	Cancelled are shown via the Status column rather than silently
	excluded, so a card that was cancelled with nothing reissued in its
	place still shows up as a real row (Status "Cancelled"), not simply
	missing - that silent-exclusion gap is exactly what looked like a
	data-loss discrepancy before this column existed."""
	if not frappe.has_permission("Fleet Card", "read"):
		frappe.throw(frappe._("Not permitted"), frappe.PermissionError)

	rows = frappe.get_all(
		"Fleet Card",
		# docstatus != 0, not docstatus == 1: a Cancelled card (docstatus 2)
		# is exactly the case the Status column exists to show - excluding
		# it here would silently reproduce the same hidden-row problem the
		# old status-based exclusion caused, just via a different filter.
		filters={"docstatus": ["!=", 0]},
		fields=[
			"registration_number",
			"allocation_type",
			"fleet_number",
			"issued_to_name",
			"issued_to",
			"bank",
			"card_number",
			"card_pin",
			"expiry_date",
			"covers_fuel",
			"covers_toll",
			"covers_oil",
			"card_limit",
			"collected",
			"status",
		],
		order_by="allocation_type asc, fleet_number asc",
	)

	return {
		"filename": f"fleet_card_assignments_{nowdate()}.xlsx",
		"content": _assignments_to_xlsx_base64(rows),
		"type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
	}


def _fleet_nr_label(row):
	"""Mirrors the source spreadsheet's own "PRIVATE <name>" convention for
	a card that isn't tied to a company Asset at all."""
	if row.allocation_type == "Asset":
		return row.fleet_number
	return f"PRIVATE - {row.issued_to_name or row.fleet_number}"


def _assignments_to_xlsx_base64(rows):
	import base64
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.worksheet.table import Table, TableStyleInfo

	wb = Workbook()
	ws = wb.active
	ws.title = "Fleet Cards"

	columns = [
		("REG NR", 16),
		("FLEET NR", 20),
		("Driver", 32),
		("Bank", 18),
		("Card Nr", 22),
		("Card Pin", 12),
		("Exp Date", 10),
		("Card Capabilities", 18),
		("Card Limit", 14),
		("Collected", 11),
		("Status", 12),
	]

	header_fill = PatternFill(fill_type="solid", fgColor="FFD9EAF7", bgColor="FFD9EAF7")
	header_font = Font(bold=True)

	for col_idx, (label, width) in enumerate(columns, start=1):
		cell = ws.cell(row=1, column=col_idx, value=label)
		cell.font = header_font
		cell.fill = header_fill
		ws.column_dimensions[cell.column_letter].width = width

	current_row = 2

	for row in rows:
		values = [
			row.registration_number or "",
			_fleet_nr_label(row),
			row.issued_to_name or row.issued_to or "",
			row.bank or "",
			row.card_number,
			row.card_pin or "",
			getdate(row.expiry_date).strftime("%m/%y") if row.expiry_date else "",
			_capabilities_label(row),
			row.card_limit or None,
			"Yes" if row.collected else "No",
			row.status or "",
		]

		for col_idx, value in enumerate(values, start=1):
			ws.cell(row=current_row, column=col_idx, value=value)

		current_row += 1

	if current_row > 2:
		# A real Excel Table needs at least one data row below the header.
		last_col_letter = ws.cell(row=1, column=len(columns)).column_letter
		table = Table(displayName="FleetCards", ref=f"A1:{last_col_letter}{current_row - 1}")
		table.tableStyleInfo = TableStyleInfo(
			name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False
		)
		ws.add_table(table)

	out = BytesIO()
	wb.save(out)
	out.seek(0)

	return base64.b64encode(out.read()).decode("utf-8")


class FleetCard(Document):
	def autoname(self):
		""""{card_number} - {issue_date}", collision-safe — not
		"{fleet_number} - {issue_date}": a Fleet Card's own identity is the
		physical card, the same reasoning behind naming Fleet Card
		Transaction after its own voucher rather than its vehicle."""
		base_name = f"{self.card_number} - {self.issue_date}"

		if not frappe.db.exists("Fleet Card", base_name):
			self.name = base_name
			return

		counter = 1

		while frappe.db.exists("Fleet Card", f"{base_name} - {counter}"):
			counter += 1

		self.name = f"{base_name} - {counter}"

	def before_submit(self):
		# Escape hatch for historical-data backfill ONLY (e.g. a card that
		# was cancelled before this system existed, with no physical signed
		# slip ever generated for it to attach) - never set by any UI path,
		# only by a one-off migration script that explicitly opts in.
		if not self.attach and not self.flags.get("ignore_attach_requirement"):
			frappe.throw(
				frappe._("Attach the signed \"Fleet Card Issued\" document before submitting."),
				title=frappe._("Attachment Required"),
			)

		if self.previous_fleet_card and not self.comments:
			frappe.throw(
				frappe._("A Reason is required when reissuing a previous Fleet Card."),
				title=frappe._("Reason Required"),
			)

		self.status = _expiry_status(self.expiry_date)

	def on_submit(self):
		if self.previous_fleet_card:
			frappe.db.set_value("Fleet Card", self.previous_fleet_card, "status", "Superseded")

	def on_cancel(self):
		# on_cancel() runs AFTER _cancel()'s own save() already wrote
		# docstatus=2 to the DB (see Document.run_post_save_methods -
		# it's called from inside _save(), post-INSERT/UPDATE) - a plain
		# `self.status = ...` here only changes the in-memory object and is
		# silently discarded, never persisted. Needs an explicit write.
		frappe.db.set_value("Fleet Card", self.name, "status", "Cancelled")

		if self.previous_fleet_card and frappe.db.get_value("Fleet Card", self.previous_fleet_card, "docstatus") == 1:
			# This card is no longer the current one for the Asset - the
			# card it had superseded becomes current again, same as the old
			# virtual status property would have concluded live (no newer
			# submitted card left to supersede it).
			previous = frappe.get_doc("Fleet Card", self.previous_fleet_card)
			frappe.db.set_value("Fleet Card", previous.name, "status", _expiry_status(previous.expiry_date))

	@property
	def days_left(self):
		if not self.expiry_date:
			return None
		return date_diff(self.expiry_date, nowdate())
